import os
import re

import config
from storage import make_object_storage, MetadataDB
import chunking as rag_chunking



def validate_upload(path: str) -> tuple[bool, str]:
    """校验文件：扩展名白名单 + 大小限制 + 魔数嗅探。返回 (ok, 错误信息)。"""
    name = os.path.basename(path)
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in config.ALLOWED_EXTS:
        return False, f"格式不允许: .{ext}，仅支持 {', '.join(sorted(config.ALLOWED_EXTS))}"
    size = os.path.getsize(path)
    if size > config.MAX_SIZE:
        return False, f"文件超过 32M 限制（{size} 字节）"
   
    with open(path, "rb") as f:
        head = f.read(16)
    if ext == "pdf" and not head.startswith(b"%PDF"):
        return False, "文件内容与 PDF 格式不符（魔数校验失败）"
    if ext == "txt" and b"\x00" in head:
        return False, "文件疑似二进制，与 txt 格式不符"
    return True, ""


def _extract_pdf(path: str) -> str:
    import pymupdf
    doc = pymupdf.open(path)
    try:
        return "\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()


def _extract_docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_html(path: str) -> str:
    from bs4 import BeautifulSoup
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f.read(), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return soup.get_text("\n")


def _extract_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extract_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            parts.append(f"## 第{i}页\n" + "\n".join(texts))
    return "\n\n".join(parts)


def _extract_ipynb(path: str) -> str:
    import json
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        nb = json.load(f)
    parts = []
    for i, cell in enumerate(nb.get("cells", [])):
        src = "".join(cell.get("source", []))
        if cell.get("cell_type") == "code":
            parts.append(f"## 代码块{i}\n```python\n{src}\n```")
        else:
            parts.append(f"## 文本块{i}\n{src}")
    return "\n\n".join(parts)


_EXTRACTORS = {
    "pdf": _extract_pdf,
    "docx": _extract_docx,
    "html": _extract_html,
    "txt": _extract_txt,
    "pptx": _extract_pptx,
    "ipynb": _extract_ipynb,
}


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lstrip(".").lower()
    fn = _EXTRACTORS.get(ext)
    if fn is None:
        raise ValueError(f"不支持的文件类型: .{ext}")
    return fn(path)


def process_upload(
    src_path: str,
    object_store=None,
    db: MetadataDB | None = None,
    vector_store=None,
) -> dict:
    """
    处理一个上传文件，返回结果信息。
    流程：校验 → 存对象存储 → 提取文本 → 智能分块 → 向量入库 → 元数据登记。
    vector_store 传入 rag.vectorstore.VectorStore 实例时执行向量化；否则只解析。
    """
    ok, err = validate_upload(src_path)
    if not ok:
        return {"ok": False, "error": err}

    filename = os.path.basename(src_path)
    ext = os.path.splitext(filename)[1].lstrip(".").lower()
    size = os.path.getsize(src_path)

    object_store = object_store or make_object_storage()
    db = db or MetadataDB()

  
    key = object_store.put(src_path)
  
    doc_id = db.create_document(filename, ext, size, key)
    try:
        db.update_status(doc_id, "parsing")
    
        text = extract_text(src_path)
       
        chunks = rag_chunking.split_text(text, source=filename)
        if not chunks:
            db.update_status(doc_id, "failed")
            return {"ok": False, "doc_id": doc_id, "error": "解析后无有效文本内容"}
      
        db.add_chunks(doc_id, chunks)
        if vector_store is not None:
            vector_store.add_chunks(chunks, doc_id)
        db.update_status(doc_id, "vectorized")
        return {
            "ok": True,
            "doc_id": doc_id,
            "filename": filename,
            "chunks": len(chunks),
            "object_key": key,
        }
    except Exception as e:  
        db.update_status(doc_id, "failed")
        return {"ok": False, "doc_id": doc_id, "error": f"{type(e).__name__}: {e}"}


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1:
        r = process_upload(sys.argv[1])
        print(r)
    else:
        print("用法: python -m phase2.modules.fileproc.parser <文件路径>")
